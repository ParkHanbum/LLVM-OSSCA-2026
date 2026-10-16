"""Convert only the 36 explicitly selected documents. Never import chat exports."""
from __future__ import annotations
import argparse, hashlib, html as H, json, logging, os, re, shutil, subprocess, sys
from pathlib import Path
from zipfile import ZipFile
from urllib.parse import urlsplit, unquote
from lxml import html, etree
from PIL import Image, ImageDraw
import pdfplumber
import pypdfium2 as pdfium
from docx import Document
from docx.text.paragraph import Paragraph
from docx.table import Table
from docx.oxml.ns import qn
from pptx import Presentation

ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / 'docs'
WORK = ROOT / '.작업자료'
logging.getLogger('pdfminer').setLevel(logging.ERROR)
ESC = H.escape
EMAIL = re.compile(r'(?<![\w@])[-\w.+]+@(?:[\w-]+\.)+[A-Za-z]{2,}(?![\w.])')
LOCAL = re.compile(r'file://[^\s<>\"\']+|/(?:Users|home)/[^\s<>\"\']+|[A-Za-z]:\\Users\\[^\s<>\"\']+')
MONO = re.compile(r'mono|consolas|courier|menlo|code', re.I)

def clean(t):
    t = EMAIL.sub('[연락처 생략]', t)
    t = LOCAL.sub('[로컬 경로 생략]', t)
    return ''.join(c for c in t if ord(c) >= 32 or c in '\n\t').replace('[이름 입력]', '이하동')

def safe_url(url):
    parts = urlsplit(url)
    return url if parts.scheme in ('http', 'https') and parts.hostname else None

def sanitize(fragment):
    """Keep semantic content. Drop scripts, event handlers and unavailable local links."""
    root = html.fragment_fromstring(fragment, create_parent='div')
    for node in root.xpath('.//script|.//style|.//iframe|.//object|.//embed|.//form|.//input|.//button|.//link|.//meta|.//base'):
        node.drop_tree()
    for node in list(root.iterdescendants()):
        if not isinstance(node.tag,str): continue
        if node.tag.lower()=='h1':node.tag='h2'
        if node.tag.lower() not in {'a','abbr','article','aside','b','blockquote','br','caption','cite','code','col','colgroup','dd','del','details','dfn','div','dl','dt','em','figcaption','figure','h1','h2','h3','h4','h5','h6','header','hr','i','img','kbd','li','main','mark','nav','ol','p','pre','s','samp','section','small','span','strong','sub','summary','sup','table','tbody','td','th','thead','tr','u','ul'}:
            node.drop_tag(); continue
        attrs = dict(node.attrib)
        for k in list(node.attrib):
            if k not in {'id','class','href','src','alt','title','colspan','rowspan','start','open'}:
                del node.attrib[k]
        if node.tag == 'a':
            href=attrs.get('href','')
            if href.startswith('#'):
                node.set('href',href)
            elif safe_url(href):
                node.set('rel','noopener noreferrer')
            else:
                node.tag='span'; node.attrib.pop('href',None)
                node.set('class','unavailable-link')
                node.set('title','원자료에 연결되어 있지만 이 자료 묶음에는 없는 파일입니다.')
                node.text=(node.text or '')
                badge=etree.SubElement(node,'small'); badge.text=' [미첨부]'
        if node.tag == 'img':
            src=attrs.get('src','')
            if not src.startswith('../assets/media/'):
                node.drop_tree(); continue
            node.set('loading','lazy'); node.set('decoding','async')
    for node in root.iter():
        if node.text: node.text=clean(node.text)
        if node.tail: node.tail=clean(node.tail)
    return ''.join(html.tostring(c,encoding='unicode') for c in root)

def save_image(blob, prefix):
    import io
    digest=hashlib.sha256(blob).hexdigest()[:14]
    rel=f'assets/media/{prefix}-{digest}.webp'
    dest=DOCS/rel; dest.parent.mkdir(parents=True,exist_ok=True)
    if not dest.exists():
        with Image.open(io.BytesIO(blob)) as im:
            im.thumbnail((2000,2400))
            im.save(dest,'WEBP',quality=91,method=5)
    return '../'+rel

def convert_html(path,sid):
    tree=html.parse(str(path))
    candidates=tree.xpath('//main') or tree.xpath('//article') or tree.xpath('//body')
    el=candidates[0]
    for nav in el.xpath('.//nav|.//*[contains(concat(" ",normalize-space(@class)," ")," toc ")]'):
        nav.drop_tree()
    el.tag='div'
    for heading in el.xpath('.//h1'):heading.tag='h2'
    return [{'n':1,'title':'본문','body':sanitize(html.tostring(el,encoding='unicode'))}]

def docx_run_html(para,doc,sid):
    out=[]
    for run in para.runs:
        t=ESC(clean(run.text))
        if run.font.name and MONO.search(run.font.name): t=f'<code>{t}</code>'
        if run.bold: t=f'<strong>{t}</strong>'
        if run.italic: t=f'<em>{t}</em>'
        out.append(t.replace('\n','<br>'))
        for blip in run._r.xpath('.//a:blip'):
            rid=blip.get(qn('r:embed'))
            if rid in doc.part.related_parts:
                src=save_image(doc.part.related_parts[rid].blob,sid)
                out.append(f'<figure><img src="{src}" alt="원자료의 설명 그림" loading="lazy"></figure>')
    return ''.join(out)

def convert_docx(path,sid):
    doc=Document(path); parts=[]; code_lines=[]
    def flush():
        if code_lines:
            parts.append('<pre><code>'+ESC(clean('\n'.join(code_lines)))+'</code></pre>');code_lines.clear()
    for block in doc.iter_inner_content():
        if isinstance(block,Paragraph):
            text=block.text; runs=block.runs
            mono=sum(len(r.text) for r in runs if r.font.name and MONO.search(r.font.name))
            if text and mono/max(1,len(text))>.72:
                code_lines.append(text); continue
            flush()
            inner=docx_run_html(block,doc,sid)
            if not inner: continue
            style=block.style.name.lower() if block.style is not None else 'normal'
            if style.startswith('heading'):
                level=int(re.search(r'\d+',style).group()) if re.search(r'\d+',style) else 2
                tag='h'+str(min(5,level+1))
            else: tag='p'
            parts.append(f'<{tag}>{inner}</{tag}>')
        elif isinstance(block,Table):
            flush(); rows=[]
            for row in block.rows:
                cells=['<td>'+''.join('<p>'+docx_run_html(pp,doc,sid)+'</p>' for pp in cell.paragraphs)+'</td>' for cell in row.cells]
                rows.append('<tr>'+''.join(cells)+'</tr>')
            parts.append('<div class="table-wrap"><table>'+''.join(rows)+'</table></div>')
    flush()
    return [{'n':1,'title':'본문','body':sanitize(''.join(parts))}]

def iter_shapes(shapes):
    for s in shapes:
        if hasattr(s,'shapes'): yield from iter_shapes(s.shapes)
        else: yield s

def shape_position(shape):
    # Some submitted PPTX producers serialize EMU positions as decimal strings.
    # Read that coordinate without changing the submitted file.
    def coordinate(name):
        try:return float(getattr(shape,name) or 0)
        except (ValueError,TypeError):
            attr='y' if name=='top' else 'x'
            vals=shape._element.xpath('.//a:xfrm/a:off/@'+attr)
            return float(vals[0]) if vals else 0
    return (round(coordinate('top')/150000),coordinate('left'))

def convert_pptx_text(path):
    prs=Presentation(path); pages=[]
    for i,slide in enumerate(prs.slides):
        blocks=[]; links=[]; texts=[]
        for shape in sorted(iter_shapes(slide.shapes),key=shape_position):
            if shape.has_table:
                rows=['<tr>'+''.join('<td>'+ESC(clean(c.text))+'</td>' for c in row.cells)+'</tr>' for row in shape.table.rows]
                blocks.append('<div class="table-wrap"><table>'+''.join(rows)+'</table></div>')
            if not shape.has_text_frame: continue
            lines=[]; mono=total=0
            for para in shape.text_frame.paragraphs:
                lines.append(para.text)
                for run in para.runs:
                    total+=len(run.text)
                    if run.font.name and MONO.search(run.font.name): mono+=len(run.text)
                    if run.hyperlink.address and safe_url(run.hyperlink.address):links.append(run.hyperlink.address)
            text='\n'.join(lines).strip()
            if not text: continue
            texts.append(text)
            if total and mono/total>.6:
                blocks.append('<pre><code>'+ESC(clean(text))+'</code></pre>')
            else:
                blocks.append('<div class="slide-text">'+''.join('<p>'+ESC(clean(t))+'</p>' for t in lines if t.strip())+'</div>')
        if slide.has_notes_slide:
            notes=slide.notes_slide.notes_text_frame.text
            if notes.strip(): blocks.append('<details><summary>발표자 노트</summary><div>'+''.join('<p>'+ESC(clean(t))+'</p>' for t in notes.splitlines() if t.strip())+'</div></details>')
        if links:
            blocks.append('<ul>'+''.join('<li><a href="'+ESC(u,quote=True)+'">'+ESC(u)+'</a></li>' for u in dict.fromkeys(links))+'</ul>')
        pages.append({'n':i+1,'title':clean(' / '.join(texts[:2]))[:110] or f'슬라이드 {i+1}','body':sanitize(''.join(blocks))})
    return pages

def redact_boxes(page):
    boxes=[]
    for w in page.extract_words(x_tolerance=2,y_tolerance=3):
        if EMAIL.search(w['text']):boxes.append((w['x0']-2,w['top']-2,w['x1']+2,w['bottom']+2))
        if 'file:///' in w['text'] or '/Users/' in w['text']:
            boxes.append((0,w['top']-2,page.width,w['bottom']+3))
    return boxes

def convert_pdf(path,sid,render=True):
    pages=[]; pdf=pdfium.PdfDocument(str(path))
    with pdfplumber.open(path) as d:
        for i,pg in enumerate(d.pages):
            text=pg.extract_text(x_tolerance=1.5,y_tolerance=3) or ''
            if '\x00' in text:
                tp=pdf[i].get_textpage(); text=tp.get_text_bounded(); tp.close()
            lines=[s.strip() for s in text.splitlines() if s.strip()]
            # Display title only. Full material is available in its rendered page.
            title=' / '.join(x for x in lines[:3] if not re.fullmatch(r'[\d\s/]+',x))[:110] or f'{i+1}쪽'
            rel=f'assets/pages/{sid}/{i+1:03d}.webp'; dest=DOCS/rel
            boxes=redact_boxes(pg)
            if render and not dest.exists():
                dest.parent.mkdir(parents=True,exist_ok=True)
                scale=min(2.5,(1800 if pg.width>pg.height else 1440)/pg.width)
                bitmap=pdf[i].render(scale=scale); im=bitmap.to_pil().convert('RGB')
                if boxes:
                    draw=ImageDraw.Draw(im)
                    for box in boxes:draw.rectangle(tuple(v*scale for v in box),fill='white')
                im.save(dest,'WEBP',quality=87,method=4)
                im.close(); bitmap.close()
            pages.append({'n':i+1,'title':clean(title),'text':clean(text),'image':'../'+rel,'ratio':round(pg.height/pg.width,6),'redacted':bool(boxes)})
            pg.close()
    pdf.close()
    return pages

def office_pdfs(paths,soffice):
    if not paths:return {}
    inp=WORK/'office-input';out=WORK/'office-pdf'
    inp.mkdir(parents=True,exist_ok=True);out.mkdir(parents=True,exist_ok=True)
    pending=[];result={}
    for sid,path in paths:
        dest=out/(sid+'.pdf');result[sid]=dest
        if dest.exists():continue
        staged=inp/(sid+'.pptx');shutil.copy2(path,staged);pending.append(staged)
    if pending:
        if not soffice:raise RuntimeError('PPTX 원본 슬라이드 변환에 LibreOffice가 필요합니다. --soffice 경로를 지정하세요.')
        env=os.environ.copy()
        if (WORK/'fonts.conf').exists():env['FONTCONFIG_FILE']=str(WORK/'fonts.conf')
        cmd=[soffice,'-env:UserInstallation=file:///private/tmp/llvm-mentee-docs-lo-20261016','--headless','--convert-to','pdf','--outdir',str(out)]+[str(p)for p in pending]
        done=subprocess.run(cmd,env=env,capture_output=True,text=True,timeout=240)
        if done.returncode:raise RuntimeError(done.stderr[-1000:])
        for sid,_ in paths:
            if not result[sid].exists():raise RuntimeError(f'{sid}: PPTX의 PDF 변환 결과가 없습니다.')
    return result

def run(args):
    source=Path(args.source).resolve(); WORK.mkdir(exist_ok=True);DOCS.mkdir(exist_ok=True)
    articles=json.loads((ROOT/'content/articles.json').read_text())
    paths=sorted({rel for a in articles for rel in a['sources']})
    if len(paths)!=36:raise RuntimeError(f'예상 문서 36개와 다릅니다: {len(paths)}')
    for rel in paths:
        p=(source/rel).resolve()
        if not p.is_relative_to(source) or not p.is_file():raise RuntimeError(f'원문 파일을 찾을 수 없습니다: {rel}')
    office=office_pdfs([(f's{i+1:03d}',source/rel)for i,rel in enumerate(paths)if Path(rel).suffix.lower()=='.pptx'],args.soffice)
    md_paths=[rel for rel in paths if Path(rel).suffix.lower()=='.md']
    payload=[(source/rel).read_text()for rel in md_paths]
    proc=subprocess.run([args.node,str(ROOT/'tools/render-markdown.cjs')],input=json.dumps(payload),capture_output=True,text=True,check=True)
    markdown=dict(zip(md_paths,json.loads(proc.stdout)))
    converted=[]
    checkpoint=WORK/'sources-cache.json'
    cached={s['path']:s for s in json.loads(checkpoint.read_text())} if checkpoint.exists() else {}
    for i,rel in enumerate(paths):
        sid=f's{i+1:03d}';path=source/rel;ext=path.suffix.lower();digest=hashlib.sha256(path.read_bytes()).hexdigest()
        article=next(a for a in articles if rel in a['sources'])
        old=cached.get(rel)
        if old and old['sha256']==digest and old['id']==sid and all((DOCS/p['image'].removeprefix('../')).is_file() for p in old['pages'] if p.get('image')):
            converted.append(old)
            print(f'{i+1:02d}/36 {sid} {rel} (cached)',flush=True)
            continue
        if ext=='.pdf':pages=convert_pdf(path,sid)
        elif ext=='.pptx':
            pages=convert_pptx_text(path);frames=convert_pdf(office[sid],sid)
            if len(pages)!=len(frames):raise RuntimeError(f'{rel}: slide/page 개수가 다릅니다.')
            for page,frame in zip(pages,frames):page.update({k:v for k,v in frame.items()if k in ['image','ratio','redacted']})
        elif ext=='.docx':pages=convert_docx(path,sid)
        elif ext=='.html':pages=convert_html(path,sid)
        else:pages=[{'n':1,'title':'본문','body':sanitize(markdown[rel])}]
        converted.append(dict(id=sid,path=rel,filename=path.name,author=article['author'],article=article['id'],title=article['title'],format=ext[1:].upper(),sha256=digest,pages=pages))
        checkpoint.write_text(json.dumps(converted,ensure_ascii=False))
        print(f'{i+1:02d}/36 {sid} {rel} ({len(pages)})',flush=True)
    (ROOT/'content/sources.json').write_text(json.dumps(converted,ensure_ascii=False,indent=2)+'\n')
    print('Source conversion complete.',flush=True)

if __name__=='__main__':
    parser=argparse.ArgumentParser()
    parser.add_argument('--source',required=True)
    parser.add_argument('--node',default='node')
    parser.add_argument('--soffice',default=shutil.which('soffice'))
    run(parser.parse_args())
